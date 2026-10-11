"""Pack the rendered slide images into docs/presentation.pptx (speaker notes from speaker_notes.md if present)."""

import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).parent
NOTES = HERE / "speaker_notes.md"
notes = dict(re.findall(r"^## (\d+)\n(.*?)(?=^## |\Z)", NOTES.read_text(), re.S | re.M)) if NOTES.exists() else {}

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
for i in range(1, 8):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    slide.shapes.add_picture(str(HERE / "png" / f"slide{i}.png"), 0, 0, prs.slide_width, prs.slide_height)
    slide.notes_slide.notes_text_frame.text = notes.get(str(i), "").strip()
prs.save(HERE.parent / "presentation.pptx")
